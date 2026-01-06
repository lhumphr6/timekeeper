import os
import io
import csv
import json
from datetime import datetime, date, timedelta
from typing import Dict, List, Optional

from fastapi import APIRouter, Request, Depends, UploadFile, File, Form, Query, HTTPException
from fastapi.responses import RedirectResponse, HTMLResponse
from sqlalchemy.orm import Session
from sqlalchemy import Column, Integer, String, DateTime, ForeignKey, func, text

from .db import get_session
from .models import Base, Employee, TimeEntry, TimesheetPeriod
from .utils import enumerate_timesheets_global

router = APIRouter(prefix="/import/department", tags=["Department Import"])

class ImportBatch(Base):
    __tablename__ = "import_batches"
    id = Column(Integer, primary_key=True, autoincrement=True)
    timesheet_id = Column(Integer, index=True, nullable=False)
    source_name = Column(String(255), nullable=True)
    created_at = Column(DateTime, default=datetime.utcnow, nullable=False)

class ImportBatchItem(Base):
    __tablename__ = "import_batch_items"
    id = Column(Integer, primary_key=True, autoincrement=True)
    batch_id = Column(Integer, ForeignKey("import_batches.id"), nullable=False, index=True)
    time_entry_id = Column(Integer, ForeignKey("time_entries.id"), nullable=False, index=True)

DATE_FMTS = ["%Y-%m-%d", "%m/%d/%Y", "%m/%d/%y"]
DT_FMTS = [
    "%Y-%m-%d %H:%M",
    "%Y-%m-%d %H:%M:%S",
    "%m/%d/%Y %H:%M",
    "%m/%d/%Y %H:%M:%S",
    "%m/%d/%y %H:%M",
    "%m/%d/%y %H:%M:%S",
]

def _parse_date(s: str) -> Optional[date]:
    s = (s or "").strip()
    if not s:
        return None
    for f in DATE_FMTS:
        try:
            return datetime.strptime(s, f).date()
        except Exception:
            continue
    for f in DT_FMTS:
        try:
            return datetime.strptime(s, f).date()
        except Exception:
            continue
    try:
        return datetime.fromisoformat(s).date()
    except Exception:
        return None

def _parse_dt(s: str) -> Optional[datetime]:
    s = (s or "").strip()
    if not s:
        return None
    for f in DT_FMTS:
        try:
            return datetime.strptime(s, f)
        except Exception:
            continue
    try:
        return datetime.fromisoformat(s)
    except Exception:
        return None

def _num(s: str) -> float:
    try:
        return float(str(s).strip())
    except Exception:
        return 0.0

def _normalize_header(h: str) -> str:
    h = (h or "").strip().lower()
    h = h.replace(" ", "").replace("_", "")
    return h

def parse_rows_from_csv(data: bytes) -> List[Dict]:
    text_stream = io.StringIO(data.decode("utf-8", errors="replace"))
    reader = csv.DictReader(text_stream)
    rows = []
    for r in reader:
        rows.append(dict(r))
    return rows

def parse_rows_from_xlsx(data: bytes) -> List[Dict]:
    from openpyxl import load_workbook
    stream = io.BytesIO(data)
    wb = load_workbook(stream, data_only=True)
    ws = wb.active
    headers = [str(ws.cell(row=1, column=c).value or "").strip() for c in range(1, ws.max_column + 1)]
    rows: List[Dict] = []
    for r in range(2, ws.max_row + 1):
        cur: Dict = {}
        for c, h in enumerate(headers, start=1):
            cur[h] = ws.cell(row=r, column=c).value
        rows.append(cur)
    return rows

def normalize_import_rows(raw_rows: List[Dict]) -> List[Dict]:
    normalized = []
    for r in raw_rows:
        lr = { _normalize_header(k): v for k, v in r.items() }
        employee_name = lr.get("employeename") or lr.get("employee") or lr.get("fullname") or lr.get("name")
        if not employee_name:
            for k in r.keys():
                if _normalize_header(k) in ("employeename", "employee", "fullname", "name"):
                    employee_name = r[k]
                    break
        date_val = lr.get("workdate") or lr.get("date")
        clock_in_val = lr.get("clockin") or lr.get("in") or lr.get("timein")
        clock_out_val = lr.get("clockout") or lr.get("out") or lr.get("timeout")
        break_val = lr.get("break") or lr.get("breakhours") or lr.get("lunch")
        pto_val = lr.get("ptohours") or lr.get("pto") or lr.get("timeoffhours")
        pto_type_val = lr.get("ptotype") or lr.get("timeofftype") or lr.get("status")
        if isinstance(date_val, datetime):
            work_date = date_val.date()
        elif isinstance(date_val, date):
            work_date = date_val
        else:
            work_date = _parse_date(str(date_val or ""))
        ci = None
        co = None
        if isinstance(clock_in_val, datetime):
            ci = clock_in_val
        elif clock_in_val not in (None, ""):
            ci = _parse_dt(str(clock_in_val))
        if isinstance(clock_out_val, datetime):
            co = clock_out_val
        elif clock_out_val not in (None, ""):
            co = _parse_dt(str(clock_out_val))
        def _combine(d: Optional[date], t: Optional[datetime]) -> Optional[datetime]:
            if not d or not t:
                return t
            return datetime(d.year, d.month, d.day, t.hour, t.minute, t.second)
        clock_in = _combine(work_date, ci) if work_date else ci
        clock_out = _combine(work_date, co) if work_date else co
        break_hours = _num(break_val or 0)
        pto_hours = _num(pto_val or 0)
        pto_type = (str(pto_type_val).strip() if pto_type_val not in (None, "") else None)
        if employee_name and work_date:
            normalized.append({
                "employee_name": str(employee_name or "").strip(),
                "work_date": work_date,
                "clock_in": clock_in,
                "clock_out": clock_out,
                "break_hours": break_hours,
                "pto_hours": pto_hours,
                "pto_type": pto_type,
            })
    return normalized

def _active_timesheet(db: Session, fallback_last=True) -> Optional[TimesheetPeriod]:
    sheets = enumerate_timesheets_global(db)
    if not sheets:
        return None
    tid = sheets[-1][0] if fallback_last else sheets[0][0]
    return db.query(TimesheetPeriod).get(tid)

def _within_period(d: date, ts: TimesheetPeriod) -> bool:
    return (d >= ts.period_start and d <= ts.period_end)


def _dedup_exists(db: Session, employee_id: int, timesheet_id: int, work_date: date, clock_in: Optional[datetime], clock_out: Optional[datetime]) -> bool:
    q = db.query(TimeEntry).filter(
        TimeEntry.employee_id == employee_id,
        TimeEntry.timesheet_id == timesheet_id,
        TimeEntry.work_date == work_date,
    )
    for r in q.all():
        if (r.clock_in or None) == (clock_in or None) and (r.clock_out or None) == (clock_out or None):
            return True
    return False

@router.get("", response_class=HTMLResponse)
def importer_home(request: Request, db: Session = Depends(get_session), timesheet_id: Optional[int] = Query(None)):
    if not request.session.get("is_admin"):
        raise HTTPException(status_code=403, detail="Admin access required")
    sheets = enumerate_timesheets_global(db)
    period_options = [{"timesheet_id": tid, "display": (name or f"{ps}..{pe}")} for tid, ps, pe, name in sheets]
    active_ts = db.query(TimesheetPeriod).get(timesheet_id) if timesheet_id else _active_timesheet(db)
    return request.app.state.templates.TemplateResponse(
        "dept_importer_upload.html",
        {"request": request, "period_options": period_options, "active_ts": active_ts.id if active_ts else None},
    )

@router.post("/upload", response_class=HTMLResponse)
async def importer_upload(
    request: Request,
    file: UploadFile = File(...),
    timesheet_id: int = Form(...),
    restrict_to_period: int = Form(1),
    db: Session = Depends(get_session),
):
    if not request.session.get("is_admin"):
        raise HTTPException(status_code=403, detail="Admin access required")
    ts = db.query(TimesheetPeriod).get(timesheet_id)
    if not ts:
        raise HTTPException(status_code=404, detail="Time Period not found")
    data = await file.read()
    ext = os.path.splitext(file.filename.lower())[1]
    if ext in (".csv", ".txt"):
        raw_rows = parse_rows_from_csv(data)
    elif ext in (".xlsx", ".xlsm"):
        raw_rows = parse_rows_from_xlsx(data)
    else:
        raise HTTPException(status_code=400, detail="Unsupported file type. Please upload CSV or XLSX.")
    norm = normalize_import_rows(raw_rows)
    if restrict_to_period:
        norm = [r for r in norm if _within_period(r["work_date"], ts)]
    by_emp: Dict[str, List[Dict]] = {}
    for r in norm:
        by_emp.setdefault(r["employee_name"], []).append(r)
    preview = []
    for name, rows in sorted(by_emp.items(), key=lambda kv: kv[0].lower()):
        emp = db.query(Employee).filter(func.lower(Employee.name) == func.lower(name)).first()
        has_any_in_period = False
        if emp:
            has_any_in_period = db.query(TimeEntry).filter(TimeEntry.timesheet_id == timesheet_id, TimeEntry.employee_id == emp.id).first() is not None
        preview.append({
            "employee_name": name,
            "status": ("Existing in period" if has_any_in_period else ("Existing employee" if emp else "New employee")),
            "existing_employee_id": emp.id if emp else None,
            "row_count": len(rows),
        })
    os.makedirs("uploads", exist_ok=True)
    slug = f"dept-import-{datetime.utcnow().strftime('%Y%m%d%H%M%S')}-{os.getpid()}.json"
    path = os.path.join("uploads", slug)
    with open(path, "w", encoding="utf-8") as f:
        json.dump({"timesheet_id": timesheet_id, "rows": norm}, f, default=str)
    return request.app.state.templates.TemplateResponse(
        "dept_importer_preview.html",
        {"request": request, "slug": slug, "timesheet_id": timesheet_id, "preview": preview},
    )

@router.post("/execute")
async def importer_execute(
    request: Request,
    slug: str = Form(...),
    timesheet_id: int = Form(...),
    selected_names: str = Form(...),
    db: Session = Depends(get_session),
):
    if not request.session.get("is_admin"):
        raise HTTPException(status_code=403, detail="Admin access required")
    path = os.path.join("uploads", slug)
    if not os.path.exists(path):
        raise HTTPException(status_code=400, detail="Import context expired. Please re-upload.")
    with open(path, "r", encoding="utf-8") as f:
        payload = json.load(f)
    if int(payload.get("timesheet_id")) != int(timesheet_id):
        raise HTTPException(status_code=400, detail="Timesheet mismatch. Please re-upload.")
    rows = payload.get("rows") or []
    def _conv_row(r: Dict) -> Dict:
        wd = r.get("work_date")
        ci = r.get("clock_in")
        co = r.get("clock_out")
        return {
            "employee_name": r.get("employee_name") or "",
            "work_date": date.fromisoformat(wd) if isinstance(wd, str) else wd,
            "clock_in": (datetime.fromisoformat(ci) if isinstance(ci, str) else ci) if ci else None,
            "clock_out": (datetime.fromisoformat(co) if isinstance(co, str) else co) if co else None,
            "break_hours": float(r.get("break_hours") or 0),
            "pto_hours": float(r.get("pto_hours") or 0),
            "pto_type": (r.get("pto_type") or None),
        }
    rows = [_conv_row(r) for r in rows]
    selected_set = {s.strip() for s in (selected_names or "").split(",") if s.strip()}
    rows = [r for r in rows if r["employee_name"] in selected_set]
    week_rows = db.execute(text("SELECT day_date, week_number FROM week_assignments WHERE timesheet_id = :tid"), {"tid": timesheet_id}).fetchall()
    week_map = {row[0]: int(row[1]) for row in week_rows}
    batch = ImportBatch(timesheet_id=timesheet_id, source_name=f"Department import {slug}", created_at=datetime.utcnow())
    db.add(batch)
    db.flush()
    inserted_count = 0
    by_emp: Dict[str, List[Dict]] = {}
    for r in rows:
        by_emp.setdefault(r["employee_name"], []).append(r)
    for name, erows in by_emp.items():
        emp = db.query(Employee).filter(func.lower(Employee.name) == func.lower(name)).first()
        if not emp:
            emp = Employee(name=name)
            db.add(emp)
            db.flush()
        for r in erows:
            wd: date = r["work_date"]
            if wd not in week_map:
                continue
            ci = r["clock_in"]
            co = r["clock_out"]
            if ci and co and co <= ci:
                co = co + timedelta(days=1)
            if _dedup_exists(db, emp.id, timesheet_id, wd, ci, co):
                continue
            total_hours = None
            if ci and co:
                total_hours = round((co - ci).total_seconds() / 3600.0, 2)
            te = TimeEntry(
                employee_id=emp.id,
                timesheet_id=timesheet_id,
                work_date=wd,
                clock_in=ci,
                clock_out=co,
                break_hours=r["break_hours"],
                total_hours=total_hours if total_hours is not None else 0,
                pto_hours=r["pto_hours"],
                pto_type=(r["pto_type"] or None),
                holiday_hours=0,
                bereavement_hours=0,
                hours_paid=0,
            )
            db.add(te)
            db.flush()
            db.add(ImportBatchItem(batch_id=batch.id, time_entry_id=te.id))
            inserted_count += 1
    db.commit()
    try:
        os.remove(path)
    except Exception:
        pass
    msg = f"Imported {inserted_count} time entries from department file."
    return RedirectResponse(url=f"/viewer?timesheet_id={timesheet_id}&msg={msg}", status_code=303)

@router.post("/undo-last")
def importer_undo_last(request: Request, timesheet_id: int = Form(...), db: Session = Depends(get_session)):
    if not request.session.get("is_admin"):
        raise HTTPException(status_code=403, detail="Admin access required")
    batch = db.query(ImportBatch).filter(ImportBatch.timesheet_id == timesheet_id).order_by(ImportBatch.created_at.desc()).first()
    if not batch:
        return RedirectResponse(url=f"/viewer?timesheet_id={timesheet_id}&msg=No+department+imports+to+undo", status_code=303)
    items = db.query(ImportBatchItem).filter(ImportBatchItem.batch_id == batch.id).all()
    ids = [it.time_entry_id for it in items]
    if ids:
        db.query(TimeEntry).filter(TimeEntry.id.in_(ids)).delete(synchronize_session=False)
    db.query(ImportBatchItem).filter(ImportBatchItem.batch_id == batch.id).delete(synchronize_session=False)
    db.delete(batch)
    db.commit()
    return RedirectResponse(url=f"/viewer?timesheet_id={timesheet_id}&msg=Undid+last+department+import", status_code=303)
